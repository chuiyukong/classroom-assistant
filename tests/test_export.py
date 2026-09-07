from io import BytesIO
import json
import zipfile

from openpyxl import load_workbook


def test_all_64_cells_and_template_parts_preserved(app, services, active):
    cls, current = active
    for n in range(1, 65):
        services[1].correct(current['id'], n, '同学' + str(n))
    output, filename = app.extensions['exports'].export(cls['id'])
    content = output.getvalue()
    cfg = json.loads((app.config['DATA_DIR'] / 'config' / 'excel-template.json').read_text(encoding='utf-8'))
    original_path = app.config['DATA_DIR'] / 'config' / 'seat-template.xlsx'
    original = load_workbook(original_path)
    exported = load_workbook(BytesIO(content))
    assert exported.sheetnames == original.sheetnames
    sheet = exported['Sheet1']
    for n in range(1, 65):
        assert sheet[cfg['seat_cells'][str(n)]].value == '同学' + str(n)
        assert sheet[cfg['number_cells'][str(n)]].value == n
    assert sheet['C3'].value == '高一（3）班  小组座位表'
    assert sheet['E18'].value == 64
    assert sheet.merged_cells == original['Sheet1'].merged_cells
    assert sheet.page_setup == original['Sheet1'].page_setup
    assert sheet.page_margins == original['Sheet1'].page_margins
    for row in original['Sheet1']:
        for cell in row:
            assert sheet[cell.coordinate].style_id == cell.style_id
    with zipfile.ZipFile(original_path) as source, zipfile.ZipFile(BytesIO(content)) as result:
        assert source.namelist() == result.namelist()
        for name in source.namelist():
            if name != 'xl/worksheets/sheet1.xml':
                assert source.read(name) == result.read(name), name
    assert filename.endswith('.xlsx')


def test_names_are_literal_text_and_empty_seats_stay_empty(app, services, active):
    values = ['=1+1', '<script>alert(1)</script>', '张&李', '+SUM(A1)', '@学生']
    for n, name in enumerate(values, 1):
        services[1].correct(active[1]['id'], n, name)
    output, _ = app.extensions['exports'].export(active[0]['id'])
    sheet = load_workbook(output)['Sheet1']
    for name, cell in zip(values, ['D15', 'D14', 'D12', 'D11', 'D9']):
        assert sheet[cell].value == name and sheet[cell].data_type == 's'
    assert sheet['X5'].value in ('', None)
    assert sheet['E18'].value == 5
